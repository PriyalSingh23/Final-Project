# -*- coding: utf-8 -*-
"""Render slides_data.py into an editable PowerPoint file.

    python presentation/build_pptx.py

Output: presentation/ANSHUMAN_deck.pptx  (16:9, dark theme, speaker notes
filled in on every slide so Presenter View works during the viva).
"""
import os
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from slides_data import S, TITLE, SUBTITLE, META                      # noqa: E402

OUT = os.path.join(HERE, "ANSHUMAN_deck.pptx")
BG = RGBColor(0x0A, 0x0E, 0x14)
INK = RGBColor(0xE6, 0xED, 0xF3)
DIM = RGBColor(0x8B, 0x98, 0xA5)
ACC = RGBColor(0x3F, 0xB9, 0x50)
ACC2 = RGBColor(0x58, 0xA6, 0xFF)
WARN = RGBColor(0xF0, 0x88, 0x3E)
LINE = RGBColor(0x1F, 0x29, 0x33)
W, H = Inches(13.333), Inches(7.5)


def blank(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    f = s.background.fill
    f.solid()
    f.fore_color.rgb = BG
    return s


def tb(slide, l, t, w, h):
    box = slide.shapes.add_textbox(l, t, w, h)
    box.text_frame.word_wrap = True
    return box.text_frame


def para(tf, text, size, color, bold=False, first=False, space_after=6, italic=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.text = text
    p.space_after = Pt(space_after)
    for r in p.runs:
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.bold = bold
        r.font.italic = italic
        r.font.name = "Segoe UI"
    return p


def title(slide, text, size=32, color=INK, top=0.55, height=1.1):
    tf = tb(slide, Inches(0.6), Inches(top), W - Inches(1.2), Inches(height))
    para(tf, text, size, color, bold=True, first=True)
    return tf


def sub(slide, text, top=1.5, size=15):
    if not text:
        return
    tf = tb(slide, Inches(0.6), Inches(top), W - Inches(1.2), Inches(0.9))
    para(tf, text.replace("`", ""), size, DIM, first=True)


def notes(slide, s):
    if not s.get("notes"):
        return
    ntf = slide.notes_slide.notes_text_frame
    ntf.text = s["notes"]
    if s.get("kid"):
        ntf.add_paragraph().text = "\n10-year-old version: " + s["kid"]


def bullets(slide, items, top=2.05, size=16):
    tf = tb(slide, Inches(0.6), Inches(top), W - Inches(1.1), H - Inches(top + 0.5))
    first = True
    for line in items:
        if line.startswith(">>"):
            para(tf, "\u25A0  " + line[2:].strip().replace("`", ""), size + 1, ACC,
                 bold=True, first=first, space_after=10)
        elif line.startswith("--"):
            para(tf, "        \u2013  " + line[2:].strip().replace("`", ""), size - 2,
                 RGBColor(0xC3, 0xCE, 0xD9), first=first, space_after=4)
        else:
            para(tf, "\u25B8  " + line.replace("`", ""), size, INK, first=first,
                 space_after=7)
        first = False


def table(slide, rows, top=2.05, size=11):
    n_r, n_c = len(rows), len(rows[0])
    shp = slide.shapes.add_table(n_r, n_c, Inches(0.55), Inches(top),
                                 W - Inches(1.1), Inches(0.34 * n_r))
    t = shp.table
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = t.cell(r, c)
            cell.text = str(val).replace("`", "")
            cell.margin_top = cell.margin_bottom = Pt(2)
            cell.margin_left = cell.margin_right = Pt(5)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            ours = r and ("ours" in str(row[0]).lower() or "this work" in str(row[0]).lower())
            cell.fill.solid()
            cell.fill.fore_color.rgb = (RGBColor(0x12, 0x1A, 0x14) if ours
                                        else (RGBColor(0x0F, 0x14, 0x1C) if r % 2
                                              else BG))
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(size)
                    run.font.name = "Segoe UI"
                    run.font.color.rgb = ACC2 if r == 0 else (ACC if ours else INK)
                    run.font.bold = (r == 0) or ours


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H

    for s in S:
        kind = s.get("kind", "bullets")
        sl = blank(prs)

        if kind == "title":
            tf = tb(sl, Inches(0.8), Inches(1.9), W - Inches(1.6), Inches(3.6))
            para(tf, "M.TECH DISSERTATION  \u00b7  DEFENCE REVIEW", 12, ACC, bold=True, first=True)
            para(tf, s["title"], 48, INK, bold=True, space_after=14)
            para(tf, s["sub"], 18, DIM, space_after=22)
            para(tf, s.get("meta", ""), 12, RGBColor(0x4B, 0x55, 0x63))

        elif kind == "section":
            tf = tb(sl, Inches(0.8), Inches(1.5), W - Inches(1.6), Inches(4))
            para(tf, str(s.get("sec", "")), 110, LINE, bold=True, first=True)
            para(tf, s["title"], 40, INK, bold=True, space_after=10)
            para(tf, s.get("sub", ""), 18, DIM)

        elif kind == "metric":
            title(sl, s["title"])
            sub(sl, s.get("sub", ""))
            cards = s["cards"]
            per = 3
            for k, (v, lab) in enumerate(cards):
                r, c = divmod(k, per)
                l = Inches(0.6 + c * 4.15)
                t = Inches(2.0 + r * 2.25)
                box = sl.shapes.add_shape(1, l, t, Inches(3.85), Inches(1.95))
                box.fill.solid()
                box.fill.fore_color.rgb = RGBColor(0x0F, 0x14, 0x1C)
                box.line.color.rgb = ACC
                box.line.width = Pt(1)
                tf = box.text_frame
                tf.word_wrap = True
                tf.margin_left = tf.margin_right = Inches(0.16)
                tf.margin_top = Inches(0.14)
                para(tf, str(v), 30, ACC, bold=True, first=True, space_after=4)
                para(tf, lab.replace("\n", " "), 11, DIM)

        elif kind == "table":
            title(sl, s["title"], size=26)
            sub(sl, s.get("sub", ""), top=1.32, size=12)
            table(sl, s["rows"], top=1.95,
                  size=13 if len(s["rows"]) <= 8 else 10)

        elif kind == "quote":
            tf = tb(sl, Inches(0.8), Inches(1.6), W - Inches(1.6), Inches(4.6))
            colors = [ACC2, ACC, DIM]
            sizes = [34, 34, 15]
            for k, chunk in enumerate(s["rows"][0]):
                para(tf, chunk, sizes[k % 3], colors[k % 3],
                     bold=k < 2, first=(k == 0), space_after=14)

        elif kind == "end":
            tf = tb(sl, Inches(0.8), Inches(2.6), W - Inches(1.6), Inches(3))
            para(tf, s["title"], 54, INK, bold=True, first=True, space_after=14)
            para(tf, s.get("sub", ""), 20, DIM)

        else:
            tag = "LIVE DEMO" if kind == "demo" else ""
            if tag:
                tf = tb(sl, Inches(0.6), Inches(0.3), Inches(3), Inches(0.4))
                para(tf, tag, 12, ACC, bold=True, first=True)
            title(sl, s["title"], size=28, top=0.62 if tag else 0.55)
            sub(sl, s.get("sub", ""), top=1.42)
            bullets(sl, s.get("bullets", []), top=1.95,
                    size=15 if len(s.get("bullets", [])) > 7 else 17)

        notes(sl, s)

    prs.save(OUT)
    print(f"[pptx] {OUT}  ({len(S)} slides, {os.path.getsize(OUT)//1024} KB)")
    print("[pptx] speaker notes are embedded on every slide \u2014 use Presenter View")


if __name__ == "__main__":
    main()
